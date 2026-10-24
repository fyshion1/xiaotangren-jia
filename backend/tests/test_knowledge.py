import asyncio
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from xiaotangren_tools import ToolContext, ToolExecutor, ToolPolicy, ToolRegistry

from app import db, plugin_state
from app.routers import knowledge as routes, personal_db
from app.services import knowledge as k, knowledge_tools, agent_tools


class KnowledgeTests(unittest.TestCase):
    def setUp(self):
        available=patch('app.personal_space.enabled',return_value=True)
        available.start();self.addCleanup(available.stop)
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.base = Path(temp.name)
        state=patch.object(plugin_state,'STATE_PATH',self.base/'plugin-state.json')
        state.start();self.addCleanup(state.stop)
        self.root = self.base / 'database'
        for obj, name, value in [(db, 'DATA_DIR', self.base / 'data'), (db, 'DB_PATH', self.base / 'app.db'),
                                 (personal_db, '_ROOT', self.root), (personal_db, '_VERSION_DIR', self.root / '.versions')]:
            mock = patch.object(obj, name, value); mock.start(); self.addCleanup(mock.stop)
        env = patch.dict(os.environ, {'PERSONAL_DATABASE_ROOT': str(self.root)})
        env.start(); self.addCleanup(env.stop)
        offline = patch.object(k.embedding, 'embed_texts', side_effect=k.embedding.EmbeddingError('offline'))
        offline.start(); self.addCleanup(offline.stop)
        db.init_db()
        db.set_settings({"expert_mode_v1": '{"enabled":true}'})
        self.china = self.make('中国历史.docx', ['唐朝', '618年', '唐朝实行三省六部制，科举制度发展。',
            '宋朝', '货币制度', '宋代出现纸币交子，商业贸易和城市经济有所发展。',
            '1000年', '此条关于纸币的来源待核实？', '明朝', '地方治理发生变化。'])
        self.japan = self.make('东亚日本史.docx', ['700年', '日本遣唐使往来和政治制度改革。', '1200年', '日本的武士政治与幕府。'])
        self.doc = k.import_history('dalaoda', [self.china])['items'][0]
        k.import_history('dalaoda', [self.japan])
        app = FastAPI(); app.include_router(routes.router, prefix='/api'); app.include_router(personal_db.router, prefix='/api')
        self.client = TestClient(app); self.addCleanup(self.client.close)

    def make(self, name, paragraphs):
        rel = k.HISTORY + '/' + name
        p = k.safe(rel); p.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        for text in paragraphs:
            doc.add_paragraph(text)
        doc.save(p)
        return rel

    def test_agent_capture_binds_current_request_without_auto_confirm(self):
        text = '```python\n    print("逐字保存")\n```\n'
        model = AsyncMock(side_effect=[{'tool_calls': [{'id': 'capture-test', 'function': {
            'name': 'knowledge_propose', 'arguments': json.dumps({
                'title': '代码片段', 'from_current_message': True, 'category': '代码'})}}]},
            {'content': '已生成待确认草稿', 'finish_reason': 'stop'}])
        config = agent_tools.ToolSettings(enabled=True, allowed_names=['knowledge_propose'])
        with patch.object(agent_tools.openai, 'complete_chat_full', model):
            result = asyncio.run(agent_tools.generate([
                {'role': 'user', 'content': '收录下面原文\n【资料原文】\n' + text}],
                {'model': 'test', 'base_url': 'http://unused'}, config,
                persona_id='dalaoda', session_id='capture-session', feature_id='core',
                request_id='capture-test', max_tokens=None))
        self.assertEqual(result.status, 'completed')
        drafts = k.captures('dalaoda')
        self.assertEqual(len(drafts), 1)
        self.assertEqual(drafts[0]['body'], text)
        self.assertFalse(k.safe(k.COLLECTED).exists())

    def test_generic_capture_confirmation_preserves_literal_text_and_is_searchable(self):
        text = '\n```python\r\n    score = 700\r\n    print("独特代码标记")\r\n```\n\n'
        draft = k.capture('dalaoda', '../代码:原文', text, '代码 / 随笔', '用户自己的片段')
        self.assertEqual(draft['status'], 'pending')
        self.assertFalse(k.safe(k.COLLECTED).exists())
        self.assertFalse(k.search('dalaoda', '独特代码标记', semantic=False)['items'])
        result = k.settle_capture('dalaoda', draft['id'], 'confirm')
        self.assertEqual(k.safe(result['source_path']).read_bytes(), text.encode())
        raw = k.safe(result['path']).read_bytes().decode()
        meta, body = k.memory_workflow.document(raw)
        self.assertEqual(body, text)
        self.assertEqual(meta['category'], '代码 / 随笔')
        hits = k.search('dalaoda', '独特代码标记', semantic=False)['items']
        self.assertTrue(hits)
        self.assertTrue(all(h['year_start'] is None and not h['region'] for h in hits))
        self.assertEqual(k.search('other', '独特代码标记', semantic=False)['items'], [])
        self.assertEqual(k.captures('dalaoda'), [])
        self.assertEqual(k.capture('dalaoda', '../代码:原文', text, '代码 / 随笔', '用户自己的片段')['id'], draft['id'])
        self.assertEqual(k.settle_capture('dalaoda', draft['id'], 'confirm')['doc_id'], result['doc_id'])
        self.assertEqual(k.safe(result['path']).read_bytes().decode(), raw)
        self.assertTrue(any(d.get('category') == '代码 / 随笔' for d in k.documents('dalaoda')))

    def test_capture_owner_rejection_limits_and_current_identity(self):
        draft = k.capture('dalaoda', '食谱', '按个人喜好调整盐量。')
        self.assertEqual(k.capture('dalaoda', '食谱', '按个人喜好调整盐量。')['id'], draft['id'])
        with self.assertRaises(KeyError):
            k.settle_capture('other', draft['id'], 'confirm')
        self.assertEqual(k.captures('other'), [])
        self.assertEqual(self.client.post('/api/knowledge/captures', json={
            'title': 'wrong', 'body': 'text', 'expected_persona': 'other'}).status_code, 409)
        self.assertEqual(k.settle_capture('dalaoda', draft['id'], 'reject')['status'], 'rejected')
        self.assertFalse(k.safe(k.COLLECTED).exists())
        self.assertNotEqual(k.capture('dalaoda', '食谱', '按个人喜好调整盐量。')['id'], draft['id'])
        for body in (' ', 'x' * (k.MAX_CAPTURE_CHARS + 1), 'hello\x00world'):
            with self.assertRaises(ValueError):
                k.capture('dalaoda', 'test', body)

    def test_capture_conflict_does_not_overwrite_existing_file(self):
        draft = k.capture('dalaoda', '资料', '原文一二三')
        target = k.safe(f"{k.WORK}/kd-{draft['id'][3:]}.md")
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(b'manual edit')
        with self.assertRaises(RuntimeError):
            k.settle_capture('dalaoda', draft['id'], 'confirm')
        self.assertEqual(target.read_bytes(), b'manual edit')
        self.assertEqual(k.captures('dalaoda')[0]['status'], 'pending')

    def test_capture_api_and_document_full_read_pagination(self):
        text = '这是资料全文的一段。\n' * 600
        r = self.client.post('/api/knowledge/captures', json={'title': '长资料', 'body': text, 'category': '随笔'})
        self.assertEqual(r.status_code, 200)
        draft = r.json()
        self.assertTrue(any(d['id'] == draft['id'] for d in self.client.get('/api/knowledge').json()['captures']))
        result = self.client.post('/api/knowledge/captures/' + draft['id'], json={'action': 'confirm'}).json()
        offset, assembled = 0, ''
        while True:
            page = k.read_page('dalaoda', result['doc_id'], offset)
            assembled += page['text']
            if page['complete']:
                self.assertIsNone(page['next_offset'])
                break
            self.assertGreater(page['next_offset'], offset)
            offset = page['next_offset']
        self.assertEqual(assembled, k.read_document('dalaoda', result['doc_id'])['content'])
        with self.assertRaises(KeyError):
            k.read_page('other', result['doc_id'])
        with self.assertRaises(ValueError):
            k.read_page('dalaoda', result['doc_id'], len(assembled) + 1)

    def test_chat_capture_copies_current_text_without_model_transcription(self):
        async def run():
            registry = ToolRegistry()
            for tool in knowledge_tools.tools(): registry.register(tool)
            executor = ToolExecutor(registry, ToolPolicy(enabled=True, allow_proposals=True,
                allowed_names={'knowledge_propose', 'knowledge_read'}))
            context = ToolContext('dalaoda', 'dalaoda', 'qq-current')
            literal = '    print("整段精确保留")\r\n\n' * 300
            with knowledge_tools.current_request([
                {'role': 'user', 'content': '别的旧内容'},
                {'role': 'user', 'content': '【当前请求】\n收录这份代码\n【资料原文】\n' + literal}]):
                response = await executor.execute('knowledge_propose', {
                    'title': '代码', 'category': '代码', 'from_current_message': True}, context)
            self.assertTrue(response.ok)
            saved = json.loads(response.content)
            draft = next(d for d in k.captures('dalaoda') if d['id'] == saved['id'])
            self.assertEqual(draft['body'], literal)
            self.assertEqual(draft['session'], 'qq-current')
            denied = await executor.execute('knowledge_propose', {'title': '代码', 'from_current_message': True}, context)
            self.assertFalse(denied.ok, 'current text must not leak to the next call')
            other = await executor.execute('knowledge_propose', {'title': '代码', 'body': 'secret'}, ToolContext('other', 'dalaoda'))
            self.assertFalse(other.ok)
        asyncio.run(run())

    def test_import_is_non_destructive_idempotent_and_retains_tables(self):
        before = k.safe(self.china).read_bytes()
        text = k.read_document('dalaoda', self.doc['id'])['content']
        self.assertIn('source-block:', text)
        self.assertEqual(k.import_history('dalaoda', [self.china])['items'][0]['status'], 'exists')
        self.assertEqual(k.safe(self.china).read_bytes(), before)
        twin = k.safe(k.HISTORY + '/副本.docx'); twin.write_bytes(before)
        self.assertEqual(k.import_history('dalaoda', [twin.relative_to(self.root).as_posix()])['items'][0]['status'], 'duplicate')
        source = self.make('欧洲史.docx', ['欧洲史的历史资料与说明文字，供表格解析测试，不是空标题。'])
        doc = Document(k.safe(source)); table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = '专有表格词条'; table.cell(0, 1).text = '独特的表格内容'
        doc.save(k.safe(source))
        imported = k.import_history('dalaoda', [source])['items'][0]
        self.assertIn('独特的表格内容', k.read_document('dalaoda', imported['id'])['content'])

    def test_lexical_fallback_temporal_region_and_unknown_dates(self):
        result = k.search('dalaoda', '宋代货币')
        self.assertEqual(result['mode'], 'keyword_fallback')
        self.assertTrue(any('交子' in r['text'] for r in result['items']))
        self.assertTrue(all(r['year_start'] is None or (r['year_end'] >= 960 and r['year_start'] <= 1279) for r in result['items']))
        japan = k.search('dalaoda', '同期日本', region='日本', year_start=618, year_end=907, semantic=False)
        self.assertTrue(any('遣唐使' in r['text'] for r in japan['items']))
        self.assertFalse(any('幕府' in r['text'] for r in japan['items']))
        self.assertEqual(k.search('other', '宋代货币', semantic=False)['items'], [])
        self.assertEqual(k.search('dalaoda', '完全不存在的知识词xyz', semantic=False)['items'], [])

    def test_dates_preserve_bce_and_no_invented_century_precision(self):
        self.assertEqual(k.year_range('约BC2070-约BC1600，470'), (-2070, -1600))
        self.assertEqual(k.year_range('约BC2070-BC1600，470'), (-2070, -1600))
        self.assertEqual(k.year_range('公元前221年 秦统一'), (-221, -221))
        self.assertEqual(k.year_range('BC2-AD2年'), (-2, 2))
        self.assertEqual(k.year_range('公元前3世纪'), (-300, -201))
        self.assertEqual(k.year_range('6-7世纪 日本出现社会危机'), (501, 700))
        self.assertEqual(k.year_range('10世纪-11世纪起 西欧城市发展'), (901, 1100))
        self.assertEqual(k.year_range('646 大化改新'), (646, 646))
        self.assertIsNone(k.year_range('1015'))
        self.assertIsNone(k.year_range('8世纪后 封君封臣制度发展'))
        self.assertIsNone(k.year_range('第1章 政治制度'))
        self.assertIsNone(k.year_range('0年'))

    def test_unrelated_paragraph_does_not_inherit_event_year(self):
        blocks = k.chunks('700年\n日本遣唐使往来。\n奥斯曼帝国扩张，时间未知。\n', '日本史')
        dated = [b for b in blocks if b['year_start'] is not None]
        self.assertTrue(any('遣唐使' in b['text'] for b in dated))
        self.assertFalse(any('奥斯曼' in b['text'] for b in dated))

    def test_proposal_confirmation_backup_reindex_and_retry(self):
        path = k.safe(self.doc['path']); source_before = k.safe(self.china).read_bytes()
        original = path.read_bytes()
        proposal = k.propose('dalaoda', self.doc['id'], '独创检索测试词：这是一项待核实的讨论。', '待核实', '宋代出现纸币交子')
        self.assertEqual(path.read_bytes(), original)
        self.assertEqual(k.search('dalaoda', '独创检索测试词', semantic=False)['items'], [])
        result = k.settle('dalaoda', proposal['id'], 'confirm')
        self.assertEqual(result['status'], 'applied')
        self.assertEqual(k.safe(result['version_path']).read_bytes(), original)
        self.assertEqual(k.safe(self.china).read_bytes(), source_before)
        self.assertTrue(k.search('dalaoda', '独创检索测试词', semantic=False)['items'])
        self.assertTrue(k.search('dalaoda', '宋代独创检索测试词', semantic=False)['items'])
        strict = k.search('dalaoda', '独创检索测试词', year_start=960, year_end=1279, semantic=False)
        self.assertFalse(any('独创检索测试词' in r['text'] for r in strict['items']))
        after = path.read_bytes(); k.settle('dalaoda', proposal['id'], 'confirm')
        self.assertEqual(path.read_bytes(), after)
        history = self.client.get('/api/personal-db/history', params={'path': self.doc['path']}).json()
        self.assertTrue(any(v['path'] == result['version_path'] for v in history))

    def test_conflict_reject_and_owner_boundaries(self):
        proposal = k.propose('dalaoda', self.doc['id'], '待核实讨论内容', '想法')
        path = k.safe(self.doc['path']); path.write_text(path.read_text(encoding='utf-8') + '\n人工修改\n', encoding='utf-8')
        before = path.read_bytes()
        response = self.client.post(f"/api/knowledge/proposals/{proposal['id']}", json={'action': 'confirm'})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(path.read_bytes(), before)
        with self.assertRaises(KeyError):
            k.settle('other', proposal['id'], 'confirm')
        self.assertEqual(k.settle('dalaoda', proposal['id'], 'reject')['status'], 'rejected')
        self.assertEqual(path.read_bytes(), before)
        with self.assertRaises(ValueError):
            k.propose('dalaoda', self.doc['id'], '内容', '想法', '不存在的原文')
        with self.assertRaises(ValueError):
            k.import_history('dalaoda', ['../escape.docx'])

    def test_external_edit_invalidates_citation_and_retrieves_new_body(self):
        hit = k.search('dalaoda', '纸币', semantic=False)['items'][0]
        path = k.safe(self.doc['path'])
        path.write_text('全新独立正文和检索标记，替换原有历史文档内容。', encoding='utf-8')
        with self.assertRaises(KeyError):
            k.passage('dalaoda', hit['id'])
        result = k.search('dalaoda', '全新独立正文', semantic=False)
        self.assertTrue(result['items'])
        self.assertFalse(any('交子' in r['text'] for r in result['items']))

    def test_api_uses_active_identity_and_escapes_citations(self):
        hit = k.search('dalaoda', '纸币', semantic=False)['items'][0]
        url = '/api/knowledge/citations/' + hit['id']
        self.assertEqual(self.client.get(url).status_code, 200)
        with patch.object(routes.db, 'get_active_persona_id', return_value='other'), patch.object(routes.db, 'get_profile', return_value={'id': 'other'}):
            self.assertEqual(self.client.get(url).status_code, 404)
            self.assertEqual(self.client.get('/api/knowledge').json()['documents'], [])
        path = k.safe(self.doc['path']); path.write_text('<script>alert(1)</script> 特殊引用检索标记。', encoding='utf-8')
        new = k.search('dalaoda', '特殊引用', semantic=False)['items'][0]
        page = self.client.get('/api/knowledge/citations/' + new['id'])
        self.assertIn('&lt;script&gt;', page.text)
        self.assertNotIn('<script>', page.text)

    def test_tool_policy_and_draft_receipt(self):
        async def check():
            registry = ToolRegistry()
            for tool in knowledge_tools.tools(): registry.register(tool)
            executor = ToolExecutor(registry, ToolPolicy(enabled=True, allow_proposals=True,
                allowed_names={'knowledge_search', 'knowledge_read', 'knowledge_propose'}))
            owner = ToolContext('dalaoda', 'dalaoda', 'single-test')
            result = await executor.execute('knowledge_search', {'query': '宋代货币'}, owner)
            self.assertTrue(result.ok)
            item = json.loads(result.content)['items'][0]
            self.assertIn('citation', item)
            result = await executor.execute('knowledge_propose', {'doc_id': self.doc['id'], 'body': '聊到的新想法，待核实。'}, owner)
            self.assertTrue(result.ok)
            self.assertEqual(json.loads(result.content)['status'], 'pending')
            denied = await executor.execute('knowledge_search', {'query': '纸币'}, ToolContext('other', 'dalaoda'))
            self.assertEqual(denied.code, 'denied')
            self.assertNotIn('knowledge_apply', registry.names())
        asyncio.run(check())

    def test_agent_loop_registers_reference_tools(self):
        cfg = agent_tools.ToolSettings(enabled=True, allowed_names=['knowledge_search'])
        model = AsyncMock(side_effect=[{'tool_calls': [{'id': 'lookup', 'function': {'name': 'knowledge_search', 'arguments': '{"query":"宋代货币"}'}}]},
                                      {'content': '笔记中提到了交子。', 'finish_reason': 'stop'}])
        with patch.object(agent_tools.openai, 'complete_chat_full', model):
            result = asyncio.run(agent_tools.generate([{'role': 'user', 'content': '按笔记讲宋代货币'}],
                {'model': 'fixture', 'base_url': 'http://unused'}, cfg, persona_id='dalaoda', session_id='single-test',
                feature_id='core', request_id='test', max_tokens=100))
        self.assertIn('交子', result.content)
        tool_messages = [m for m in model.await_args_list[-1].args[0] if m['role'] == 'tool']
        self.assertTrue(any('交子' in m['content'] for m in tool_messages))

    def test_editor_hash_supports_repeated_saves_and_index_refresh(self):
        path = self.doc['path']
        initial = self.client.get('/api/personal-db/file', params={'path': path}).json()
        first = self.client.put('/api/personal-db/file', json={'path': path, 'content': initial['content']+'\n增补甲\n', 'expected_sha256': initial['sha256']})
        self.assertEqual(first.status_code, 200)
        second = self.client.put('/api/personal-db/file', json={'path': path, 'content': initial['content']+'\n编辑独有乙检索\n', 'expected_sha256': first.json()['sha256']})
        self.assertEqual(second.status_code, 200)
        self.assertTrue(k.search('dalaoda', '编辑独有乙检索', semantic=False)['items'])

    def test_semantic_index_and_unchanged_chunk_vector_reuse(self):
        with patch.object(k.embedding, 'embed_texts', side_effect=lambda texts, **kw: [[1.0, 0.0] for _ in texts]):
            progress = k.index_batch('dalaoda')
            self.assertEqual(progress['remaining'], 0)
            self.assertGreater(progress['done'], 0)
            result = k.search('dalaoda', '纸币')
            self.assertEqual(result['mode'], 'hybrid')
        with k.store() as c:
            before = c.execute('SELECT count(*) FROM knowledge_chunks WHERE vector IS NOT NULL').fetchone()[0]
        p = k.propose('dalaoda', self.doc['id'], '新增资料，尚待核实。', '待核实')
        k.settle('dalaoda', p['id'], 'confirm')
        with k.store() as c:
            after = c.execute('SELECT count(*) FROM knowledge_chunks WHERE vector IS NOT NULL').fetchone()[0]
        self.assertEqual(before, after)


if __name__ == '__main__':
    unittest.main()
